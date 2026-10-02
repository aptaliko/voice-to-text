"""Turn the edited text into a Word document."""

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def build_docx(text: str, title: str = "") -> bytes:
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    if title.strip():
        heading = document.add_paragraph()
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = heading.add_run(title.strip())
        run.bold = True
        run.font.size = Pt(14)

    # Blank lines separate paragraphs; single newlines are line breaks.
    for block in text.replace("\r", "").split("\n\n"):
        block = block.strip("\n")
        if not block.strip():
            continue
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for index, line in enumerate(block.split("\n")):
            if index:
                paragraph.add_run().add_break()
            paragraph.add_run(line)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
